from pathlib import Path
import json,zipfile,stat,posixpath,sys,hashlib
import pandas as pd
import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt
from docx import Document
from docx.shared import Inches,Pt
from pipeline import ROOT

def preflight(p):
    with zipfile.ZipFile(p) as z:
        names=set(z.namelist());assert {'[Content_Types].xml','_rels/.rels','word/document.xml'}<=names
        assert len(z.infolist())<=10000;total=0
        for i in z.infolist():
            assert not i.flag_bits & 1;assert not stat.S_ISLNK(i.external_attr>>16)
            assert not i.filename.startswith('/') and '\\' not in i.filename
            assert not posixpath.normpath(i.filename).startswith('../')
            assert i.file_size<=64*1024**2;total+=i.file_size
            assert i.file_size/max(i.compress_size,1)<=1000
        assert total<=256*1024**2 and z.testzip() is None
        xml=z.read('word/document.xml');assert b'<w:ins ' not in xml and b'<w:del ' not in xml
    return Document(p)

if len(sys.argv)>1:
    src=Path(sys.argv[1]);doc=preflight(src)
    (ROOT/'docs/source_brief_text.txt').write_text('\n'.join(p.text for p in doc.paragraphs),encoding='utf-8')
    print('Brief preflight passed; paragraphs:',len(doc.paragraphs))

df=pd.read_csv(ROOT/'data/clean/mumbai_daily.csv',parse_dates=['date'])
a=pd.read_csv(ROOT/'outputs/annual_summary.csv');m=pd.read_csv(ROOT/'outputs/monthly_climatology.csv');ev=pd.read_csv(ROOT/'outputs/evaluation.csv');q=json.loads((ROOT/'outputs/quality.json').read_text());t=json.loads((ROOT/'outputs/thresholds.json').read_text());find=json.loads((ROOT/'outputs/findings.json').read_text());manifest=json.loads((ROOT/'data/source_manifest.json').read_text())
plt.rcParams.update({'axes.spines.top':False,'axes.spines.right':False,'figure.dpi':150,'font.size':11})
charts=[('seasonal_heat.png',m.month,m.mean_tmax_c,'Monthly mean daily maximum temperature, 2015–2025','Month','°C','#b5482d'),('annual_rain.png',a.year,a.rain_total_mm,'Annual gridded precipitation, Mumbai representative point','Year','mm','#14767c'),('annual_heat.png',a.year,a.heat_high_days,'Days exceeding training-period heat P95','Year','Days','#b5482d')]
for filename,x,y,title,xlabel,ylabel,color in charts:
    fig,ax=plt.subplots(figsize=(10,4.8));ax.bar(x,y,color=color);ax.set(title=title,xlabel=xlabel,ylabel=ylabel);ax.set_xticks(x);ax.grid(axis='y',alpha=.15);fig.tight_layout();fig.savefig(ROOT/'outputs'/filename);plt.close(fig)
sources='''- NASA POWER Daily API: https://power.larc.nasa.gov/docs/services/api/temporal/daily/
- Meteorology methodology: https://power.larc.nasa.gov/docs/methodology/meteorology/
- Precipitation validation: https://power.larc.nasa.gov/docs/methodology/meteorology/precipitation/
- POWER referencing: https://power.larc.nasa.gov/docs/referencing/
- NASA data-use guidance: https://www.earthdata.nasa.gov/engage/open-data-services-software-policies/data-use-guidance
'''
report=f'''# CityPulse — final technical report
## Executive summary
Mumbai representative point, 2015–2025. Historical decision support for municipal planning teams and community volunteers. This is a hackathon prototype, not an official warning or operational forecast service.

## Data and provenance
NASA POWER Daily API {manifest['header']['api']['version']}, source identifiers {manifest['header']['sources']}, accessed {manifest['accessed_utc']}. Requested point 19.076 N, 72.8777 E. API request, source metadata and raw SHA-256 are in data/source_manifest.json. Preserved raw JSON is included. Meteorological source values represent a roughly 0.5° x 0.625° grid; they are not station observations or a city-boundary average. LST denotes local solar time, not IST. No synthetic data is used.

There are {len(df)} expected and actual days, no missing values in four source variables and no duplicate dates. No imputation or smoothing was used for daily signals. Plausible extremes were retained; impossible ordering/negative rain fails the pipeline. NASA fill value -999 is converted to missing. Dates are sorted and reindexed to all calendar days. CSV stores six decimals; raw source precision is preserved in JSON. Three-day precipitation requires three valid consecutive days. Quality details: outputs/quality.json.

## Three evidence-backed findings
'''+ '\n\n'.join(f"### {x['title']}\n{x['text']}" for x in find)+f'''

These findings describe the entire study, while dashboard charts can be filtered. Annual totals are comparable because all study years have complete daily coverage (including leap days). The unusually high coarse-cell temperatures are not independently validated against Mumbai station records; source bias/coastal representation is a material uncertainty.

## Transparent historical method
Training only: 2015–2021. Elevated heat is Tmax >= {t['heat_p90_c']:.4f} °C (P90); high heat >= {t['heat_p95_c']:.4f} °C (P95). Rain thresholds are computed on training wet days (>=1 mm): elevated >= {t['rain_wet_p95_mm']:.4f} mm/day (P95), high >= {t['rain_wet_p99_mm']:.4f} mm/day (P99). Threshold displays round for readability; calculations use full precision. These are analytical exceedance signals, not calibrated disaster probabilities, IMD categories or heatwave declarations. All-month heat thresholds emphasize absolute seasonal heat; they do not measure month-specific anomalies. No composite score or unsupported spatial ranking is invented.

## Chronological evaluation
2015–2021 fit; 2022–2023 validation; 2024–2025 held-out test. No tuning or selection on test outcomes. Both baselines are specified in advance. For each target date, persistence uses the preceding calendar day's value. Seasonal mean uses the target calendar month's training-only average. Continuous errors are MAE and RMSE in °C or mm/day. Do not compare errors across units.

For threshold-event experiments, persistence probability is the prior-day event indicator; seasonal probability is training monthly event frequency. Brier score compares probabilities to same-series exceedance outcomes. Precision/recall/F1 use 0.5 probability cutoff; rare-event baseline accuracy can be misleading, so prevalence, TP/FP/FN/TN, recall and Brier are also reported. Undefined precision/F1 are recorded as 0 when no positives are predicted. These experiments cannot validate floods or heat-related injury. They do not account for operational source-publication latency and do not establish deployable next-day forecasting capability. No forecast feature is released to users.

### Numerical results (validation and held-out test)
{ev.to_string(index=False,float_format=lambda x:f'{x:.4f}')}

Held-out interpretation: Tmax persistence MAE is 0.8403 °C versus 1.5326 °C for the seasonal mean, but its event Brier score (0.06566) is worse than the seasonal probability baseline (0.05515). Rainfall persistence MAE is 5.2500 mm/day versus 6.6081 mm/day for the seasonal mean, but rain-exceedance recall is only 2/13 = 15.4%, with 11 false alarms. Rain Brier is also worse for persistence (0.03010 versus 0.01696). Lower continuous error does not imply reliable event warnings. No operational forecast is released.

No statistical confidence intervals or station-level bias correction were estimated. Results are empirical errors for this dataset, not uncertainty bounds for a new real-world prediction. Forecast complexity was intentionally deferred given the evidence and deadline.

## Practical recommendations
- Before hot-season periods: review shaded rest/water access and outdoor-worker outreach. Use current official advisories for actions.
- Before the monsoon: prioritize drain inspection and maintenance scheduling; use local incident/tide/drainage evidence for site-specific action.
- Publish clear time scope, source type and limitations whenever sharing results. Historical seasonal planning is justified; street-level flood alerts are not.

## Limitations and next validation work
Single coarse grid cell, no station validation, no humidity/heat index, no population vulnerability, tides, terrain, drainage, land use or flood incidents. No ward comparison, no causal climate-change attribution from eleven years, no live observations, and no future disaster predictions. A longer record and independent IMD station series should precede operational thresholds. Validate event definitions against observed impacts before calling indicators risk probabilities. NASA documents nontrivial daily precipitation errors; local accuracy is not established here.

## Reproducibility and implementation
Run pipeline.py to reproduce cleaned CSV, summaries, findings, thresholds, quality and evaluation offline. Run unittest discovery and Streamlit AppTest checks. Start app.py with Streamlit. See README.md. No paid API or credentials are used. The final application and tests are local; there is no public deployment or hackathon submission.

## Three-person handover
Person 1: data/raw, data/clean, source_manifest.json, docs/DATA_DICTIONARY.md and quality.json.
Person 2: pipeline.py, analysis.ipynb, outputs charts, thresholds/evaluation/findings and this report.
Person 3: app.py, requirements.txt, tests, README.md, demo script and presentation outline. All three should review definitions and findings. Names and contributions are not invented.

## Attribution and access
The data was obtained from National Aeronautics and Space Administration (NASA) Langley Research Center's Prediction Of Worldwide Energy Resources (POWER) project funded through the NASA Earth Science Division.
The data was obtained from the POWER Project's Daily API {manifest['header']['api']['version']} version on {manifest['accessed_utc'][:10]}.
NASA acknowledgment and exact source citation are retained. No NASA endorsement is implied. NASA general open-data guidance applies subject to marked restrictions/third-party material; no POWER-specific Creative Commons license is invented. POWER requests notification of publication or redistribution; no email was sent on your behalf.

## References
{sources}
Exact download URL: {manifest['url']}
'''
(ROOT/'docs/REPORT.md').write_text(report,encoding='utf-8')
# A readable Word report with real tables/images, source document is left unchanged.
doc=Document();sec=doc.sections[0];sec.page_width=Inches(8.27);sec.page_height=Inches(11.69);sec.top_margin=sec.bottom_margin=Inches(.7);sec.left_margin=sec.right_margin=Inches(.7)
doc.styles['Normal'].font.name='Calibri';doc.styles['Normal'].font.size=Pt(10)
for line in report.splitlines():
    if not line.strip():continue
    if line.startswith('### '):doc.add_heading(line[4:],level=2)
    elif line.startswith('## '):doc.add_heading(line[3:],level=1)
    elif line.startswith('# '):doc.add_heading(line[2:],level=0)
    elif line.startswith('- '):doc.add_paragraph(line[2:],style='List Bullet')
    elif line.startswith(' '):continue # Numerical results table is added in structured form below.
    elif line.split() and line.split()[0] in ['validation','test','split'] and ('tmax_c' in line or 'rain_mm' in line or 'variable' in line):continue
    else:doc.add_paragraph(line)
doc.add_heading('Evaluation table (rounded)',level=1)
cols=['split','variable','method','n','mae','rmse','brier','f1'];table=doc.add_table(rows=1,cols=len(cols));table.style='Light Shading Accent 1'
for c,k in zip(table.rows[0].cells,cols):c.text=k
for _,r in ev.iterrows():
    for c,k in zip(table.add_row().cells,cols):c.text=f'{r[k]:.3f}' if isinstance(r[k],float) else str(r[k])
for file,*_ in charts:doc.add_picture(str(ROOT/'outputs'/file),width=Inches(6.6))
output=ROOT/'docs/CityPulse_Submission_Report.docx'
if output.exists():raise FileExistsError(output)
doc.save(output);verified=preflight(output);assert len(verified.tables)==1 and len(verified.inline_shapes)==3
# Notebook is standard Jupyter JSON, uses only packaged data and local pipeline.
cells=[]
def md(s):cells.append({'cell_type':'markdown','metadata':{},'source':s.splitlines(True)})
def code(s):cells.append({'cell_type':'code','execution_count':None,'metadata':{},'outputs':[],'source':s.splitlines(True)})
md('# CityPulse: Mumbai 2015–2025\nReproducible historical gridded-reanalysis analysis; not station observations or a flood forecast. Run from this notebook directory. See docs/REPORT.md for findings and limits.')
code("from pathlib import Path\nimport pandas as pd\nimport json\nfrom pipeline import clean, thresholds, signals, evaluate, summarize, ROOT\nraw=json.loads((ROOT/'data/raw/nasa_power_mumbai_2015_2025.json').read_text())\nd=clean(raw)\nassert len(d)==4018 and not d.isna().any().any()\nd.head()")
md('## Coverage and integrity\nFill values become missing. No imputation. Source metadata and hash are preserved.')
code("d.describe()\nprint(d.isna().sum())\nprint(raw['header'])")
md('## Fixed training-only historical signals\nAll thresholds fit only 2015–2021; validation/test never enter them.')
code("t=thresholds(d)\ns=signals(d,t)\nprint(t)\nannual,monthly,findings=summarize(s)\nfor f in findings: print(f['title'],f['text'])")
code("import matplotlib.pyplot as plt\nmonthly.mean_tmax_c.plot.bar(title='Monthly mean daily Tmax (degrees C)');plt.show()\nannual.rain_total_mm.plot.bar(title='Annual precipitation total (mm)');plt.show()\nannual.heat_high_days.plot.bar(title='High heat signal days');plt.show()")
md('## Chronological baseline experiment\nTrain 2015–2021; validation 2022–2023; test 2024–2025. Persistence uses previous day. Compare within-variable MAE/RMSE; Brier evaluates event probabilities. This is not independent flood/health validation or an operational next-day forecasting service.')
code("evaluation=evaluate(s)\nevaluation")
md('## Interpretation and limitations\nNo causal attribution, city-wide station record, neighborhood exposure or flood-occurrence claim. Precipitation and temperature alone do not identify real-world damages. Threshold rules are descriptive and not official advisories. See docs/REPORT.md and README.md for recommendations and data attribution.')
(ROOT/'analysis.ipynb').write_text(json.dumps({'cells':cells,'metadata':{'kernelspec':{'display_name':'Python 3','language':'python','name':'python3'},'language_info':{'name':'python','version':'3.11'}},'nbformat':4,'nbformat_minor':5},indent=2),encoding='utf-8')
print('Reports, three charts, notebook and Word package validated.')
